import sys, re, copy, os, subprocess, base64, docx
sys.path.insert(0, "."); from dlib import *; from dlib import _mkrun, SCRIPT_SPLIT, set_bidi; from refdata import *
from build_cert import replace_text
from docx.oxml.ns import qn
from docx.shared import Pt, Mm
sys.path.insert(0, B); os.environ.setdefault("B", B)
import make_templates2 as mt
CHROME = [os.path.join(dp, f) for dp, dn, fn in os.walk("/opt/pw-browsers/chromium-1194") for f in fn if f == "chrome"][0]
LOGO = B + "/assets/logo_cert.png"; LOGO64 = base64.b64encode(open(LOGO, "rb").read()).decode()
def header_jpg(en, ar, out):
    html = f"""<html><head><meta charset='utf-8'><style>body{{margin:0;background:#fff;width:2000px;height:150px;display:flex;align-items:center;font-family:Arial,'DejaVu Sans',sans-serif}}
.l{{flex:0 0 700px;padding-left:10px;color:#1f3864;font-weight:bold;font-size:42px;line-height:1.15}}.m{{flex:1;text-align:right;direction:rtl;color:#1f3864;font-weight:bold;font-size:48px;line-height:1.1;padding-right:30px}}
.r{{flex:0 0 420px;text-align:right}}.r img{{height:120px;width:auto}}</style></head><body><div class='l'>{en}<br>Application Form</div><div class='m'>{ar}<br>نموذج الطلب</div><div class='r'><img src='data:image/png;base64,{LOGO64}'></div></body></html>"""
    out = os.path.abspath(out); h = out + ".html"; open(h, "w").write(html)
    subprocess.run([CHROME, "--headless", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--window-size=2000,400", f"--screenshot={out}.png", "file://" + h], capture_output=True)
    subprocess.run(["convert", out + ".png", "-crop", "2000x150+0+0", "+repage", "-background", "white", "-alpha", "remove", "-quality", "92", out], capture_output=True)
AR_NAMES = {("FP001", "FCC"): ("Female Cancer Cover / Second Medical Opinion", "سرطانات الإناث / الاستشارة الطبية الثانية"),
            ("HC001", "HCB"): ("Hospital Cash Benefit due to any cause (per day)", "منفعة الاستشفاء النقدية لأي سبب (يوميًا)"),
            ("CI001", "CIB"): ("Critical Illness Cover / Second Medical Opinion", "تغطية الأمراض الحرجة / الاستشارة الطبية الثانية")}
def name_of(code, c): return BT.get((code, c)) or AR_NAMES[(code, c)]
CLAUSE_KEYS = ("Agreement", "Authorization", "Review Period", "Cancellation", "Exclusions", "Disclaimer")
def terms_for(code):
    app = PRODUCT[code]["app"] or PRODUCT["TL001"]["app"]
    out = {}; quest = []
    for en, ar in mt.docx_terms(app):
        key = next((k for k in CLAUSE_KEYS if re.match(r"^(\d-\s*)?" + k, en)), None)
        if key: out[key] = (en, ar)
        elif not en.startswith("OMR 500"): quest.append((en, ar))
    return out, quest
def split_label(txt):
    m = re.match(r"^(\d-\s*)?([A-Za-z ]+?)\s*:\s*(.*)$", txt, re.S)
    if m: return (m.group(1) or "") + m.group(2) + " : ", m.group(3)
    m = re.match(r"^([؀-ۿ ]+?):\s*(.*)$", txt, re.S)
    return (m.group(1) + ": ", m.group(2)) if m else ("", txt)
def set_labeled(tc, label, body):
    """Cell with a bold label followed by regular text (keeps the two formats of the reference cell)."""
    runs = list(tc.iter(qn("w:r"))); rp_lab = runs[0].find(qn("w:rPr")) if runs else None
    rp_body = None
    for r in runs[1:]:
        rp = r.find(qn("w:rPr"))
        if rp is not None and rp.find(qn("w:b")) is None: rp_body = rp; break
    if rp_body is None and rp_lab is not None:
        rp_body = copy.deepcopy(rp_lab)
        for e in rp_body.findall(qn("w:b")): rp_body.remove(e)
    ps = tc.findall(qn("w:p")); p = ps[0]
    for x in ps[1:]: tc.remove(x)
    for ch in list(p):
        if ch.tag != qn("w:pPr"): p.remove(ch)
    for txt, rp in ((label, rp_lab), (body, rp_body)):
        if not txt: continue
        for sub in SCRIPT_SPLIT.split(txt):
            if sub: p.append(_mkrun(sub, rp, rtl=bool(AR.search(sub))))
def build_app(code, sub, out, workdir="work"):
    os.makedirs(workdir, exist_ok=True); pinfo = PRODUCT[code]
    d = docx.Document(REF_APP); T = d.tables; body = d.element.body
    # ---- header image (page 1 and page 2 share one picture)
    hj = f"{workdir}/hdr_{code}.jpg"; header_jpg(pinfo["en"], pinfo["ar"], hj)
    for rel in d.part.rels.values():
        if rel.reltype.endswith("/image") and rel.target_part.partname.endswith("image1.jpeg"): rel.target_part._blob = open(hj, "rb").read()
    # ---- applicant table
    t0 = T[0]; R = lambda i: tcs(t0.rows[i]); sv = lambda f: tag(f"{RISK}.{f}", f)
    set_tc(R(0)[1], tag("PolicyObject.ProposalNo", "ProposalNo")); set_tc(R(0)[2], "رقم نموذج الطلب")
    set_tc(R(1)[1], tag("PolicyObject.ApplyDate", "ApplyDate")); set_tc(R(1)[2], "تاريخ نموذج الطلب")
    set_tc(R(2)[0], f"Customer Name ({pinfo['insured']})"); set_tc(R(2)[1], sv("InsuredName")); set_tc(R(2)[2], "اسم العميل (المؤمن عليه)")
    set_tc(R(4)[1], sv("ResidentStatus_CodeDesc")); set_tc(R(5)[0], "Passport / ID Number"); set_tc(R(5)[2], "رقم جواز السفر / البطاقة الشخصية")
    set_tc(R(6)[2], "عنوان المراسلات"); set_tc(R(7)[2], "رقم الهاتف النقال"); set_tc(R(8)[2], "عنوان البريد الإلكتروني")
    # ---- beneficiary table
    t1 = T[1]
    if pinfo["app_benef"]:
        h = tcs(t1.rows[0]); set_tc(h[2], "نسبة المساهمة (بمجموع يساوي 100%)"); set_tc(h[3], "Relationship to the Insured Person", "صلته بالشخص المصاب")
    else:
        prev = t1._tbl.getprevious()
        if prev is not None and prev.tag == qn("w:p"): body.remove(prev)
        body.remove(t1._tbl)
    # ---- benefits table
    t2 = T[2]; rows = list(t2.rows)
    set_tc(tcs(rows[0])[1], "المنافع ومبلغ التأمين"); set_tc(tcs(rows[1])[0], "Benefit", "المنفعة"); set_tc(tcs(rows[1])[1], "Sum Insured (OMR)", "مبلغ التغطية (بالريال العُماني)")
    # header second cell holds both labels in the reference; keep it simple: EN+AR in the SI cell
    proto = copy.deepcopy(rows[2]._tr)
    for r in rows[2:]: delete_row(r)
    anchor = rows[1]._tr; cov = COV[(code, sub)]; idx = {c: i for i, c in enumerate(cov)}
    def add_row(en, ar, si):
        nonlocal anchor
        n = copy.deepcopy(proto); anchor.addnext(n); anchor = n; c = n.findall(qn("w:tc"))
        set_tc(c[0], en, ar); set_tc(c[1], si); set_tc(c[2], si)
    pn = tag(f"{PLAN}.PlanName", "PlanName")
    add_row("Plan Opted", "البرنامج المختار", pn)
    show = SHOW.get((code, sub)) or cov
    for c in show:
        en, ar = name_of(code, c)
        if code == "CI001": en = en.replace("Cover", f"Cover ({CI_COUNT[sub]})")
        add_row(en, ar, "Included / مشمول" if c == "TRV" else cov_tag(idx[c]))
    # ---- payment table
    t3 = T[3]; pr = list(t3.rows)
    c0 = tcs(pr[0])[0]; replace_text(c0, "Annual Premium", tag("PolicyObject.PremiumModeCode_CodeDesc", "PremiumModeCode_CodeDesc"))
    set_tc(tcs(pr[1])[0], "Premium (OMR)"); set_tc(tcs(pr[1])[1], sv("DuePremium")); set_tc(tcs(pr[1])[2], "القسط (ريال عُماني)")
    if code == "LP001":
        set_tc(tcs(pr[2])[0], mt.PERIOD_LP_EN.replace("<br>", " ")); set_tc(tcs(pr[2])[2], mt.PERIOD_LP_AR.replace("<br>", " "))
    set_tc(tcs(pr[3])[0], "Customer Signature")
    # ---- declarations table
    terms, quest = terms_for(code); t4 = T[-1] if len(T) == 4 else T[3 if not pinfo["app_benef"] else 4]
    t4 = [t for t in d.tables if len(t.columns) == 26][0]; dr = list(t4.rows)
    for ri, key in ((6, "Exclusions"), (7, "Disclaimer")):
        en, ar = terms.get(key) or terms_for("TL001")[0][key]
        if code == "LP001": en = en.replace("Term Life Insurance", "Life Protect"); ar = ar.replace("التأمين على الحياة لأجل محدد", "خطة حماية الحياة")
        tc_ = tcs(dr[ri]); le, be = split_label(en.replace("\xa0", " ")); la, ba = split_label(ar.replace("\n", " "))
        set_labeled(tc_[1], le, be); set_labeled(tc_[0], la, ba)
    # ---- questionnaire (Term Life / Critical Illness application forms)
    if quest and code in ("TL001", "CI001", "LP001") and code != "LP001":
        p = d.add_paragraph(); p.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        h = d.add_paragraph(); r = h.add_run("Health and Lifestyle Questionnaire / الاستبيان الصحي ونمط الحياة"); r.bold = True; r.font.size = Pt(11)
        tb = d.add_table(rows=0, cols=2)
        bd = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            x = OxmlElement("w:" + e); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:space"), "0"); x.set(qn("w:color"), "000000"); bd.append(x)
        tb._tbl.tblPr.append(bd)
        for en, ar in quest:
            row = tb.add_row(); a, b = row.cells; a.text = en; b.text = ar
            for cell, rtl in ((a, False), (b, True)):
                for pp in cell.paragraphs:
                    for rr in pp.runs: rr.font.size = Pt(8.5)
                    if rtl: set_bidi(pp._p, True)
    for h in list(d.element.iter(qn("w:highlight"))): h.getparent().remove(h)
    d.save(out)
if __name__ == "__main__": build_app(sys.argv[1], sys.argv[2], sys.argv[3])
