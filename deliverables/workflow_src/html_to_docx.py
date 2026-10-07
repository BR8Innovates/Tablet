#!/usr/bin/env python3
"""Converts the Output Service HTML templates (tpl2/) to Word. usage: html_to_docx.py in.html out.docx
Only the constructs the templates use are handled: tables (colgroup widths, colspan, borders, shading, min height), divs, b/span/br, base64 images, RTL Arabic cells."""
import re, sys, base64, io, html
from html.parser import HTMLParser
from docx import Document
from docx.shared import Mm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
VOID = {"img", "br", "col", "meta", "link", "hr"}
class N:
    def __init__(s, tag, attrs): s.tag = tag; s.a = dict(attrs); s.c = []
    cls = property(lambda s: (s.a.get("class") or "").split())
    st = property(lambda s: dict((k.strip(), v.strip()) for k, v in (x.split(":", 1) for x in (s.a.get("style") or "").split(";") if ":" in x)))
class P(HTMLParser):
    def __init__(s): super().__init__(convert_charrefs=True); s.root = N("root", []); s.stack = [s.root]; s.skip = 0
    def handle_starttag(s, tag, attrs):
        if tag in ("style", "title", "head"): s.skip += 1; return
        if tag in ("html", "body", "meta", "tbody"): return
        n = N(tag, attrs); s.stack[-1].c.append(n)
        if tag not in VOID: s.stack.append(n)
    def handle_endtag(s, tag):
        if tag in ("style", "title", "head"): s.skip -= 1; return
        if tag in ("html", "body", "tbody") or tag in VOID: return
        for i in range(len(s.stack) - 1, 0, -1):
            if s.stack[i].tag == tag: del s.stack[i:]; return
    def handle_data(s, d):
        if not s.skip and d: s.stack[-1].c.append(d)
doc_kind = "app"
def px(v, base=None):
    m = re.match(r"([\d.]+)(mm|pt|%)?", v or ""); return (float(m.group(1)), m.group(2)) if m else (None, None)
def set_rtl_run(run, font):
    rpr = run._r.get_or_add_rPr(); rpr.append(OxmlElement("w:rtl"))
    rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(k), font)
def fmt_run(run, font, size, bold=False, color=None, rtl=False):
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
    rpr = run._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), font); rf.set(qn("w:cs"), font)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if rtl: set_rtl_run(run, "Segoe UI")
def para_fmt(p, rtl=False, align=None, before=0, after=0):
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.0
    ppr = p._p.get_or_add_pPr()
    if rtl: ppr.append(OxmlElement("w:bidi"))
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        if rtl: jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "left"); ppr.append(jc)   # start edge of a bidi paragraph = right
        else: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left" and rtl: jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "right"); ppr.append(jc)
def add_img(p, node, width_mm=None, height_mm=None):
    src = node.a.get("src", ""); m = re.match(r"data:[^;]+;base64,(.*)", src, re.S)
    if not m: return
    data = base64.b64decode(m.group(1))
    if data[:2] == b"\xff\xd8" and data[2:4] == b"\xff\xed":   # Photoshop-header JPEG: python-docx cannot read it, re-encode as plain JFIF
        import subprocess
        data = subprocess.run(["convert", "jpg:-", "-strip", "-interlace", "none", "jpg:-"], input=data, capture_output=True).stdout
    w = node.st.get("width"); wmm = px(w)[0] if w and px(w)[1] == "mm" else None
    run = p.add_run()
    if height_mm: run.add_picture(io.BytesIO(data), height=Mm(height_mm))
    else: run.add_picture(io.BytesIO(data), width=Mm(width_mm or wmm or 60))
    return run
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)
def borders(tbl, on):
    tblPr = tbl._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement("w:" + e); x.set(qn("w:val"), "single" if on else "nil"); x.set(qn("w:sz"), "4"); x.set(qn("w:space"), "0"); x.set(qn("w:color"), "000000"); b.append(x)
    tblPr.append(b)
def cell_margins(cell, top, side):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", side), ("bottom", top), ("right", side)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(int(v * 56.7))); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)
def text_of(n): return n if isinstance(n, str) else "".join(text_of(x) for x in n.c)
def inline(p, nodes, font, size, bold, rtl, color=None):
    for n in nodes:
        if isinstance(n, str):
            t = re.sub(r"\s+", " ", n.replace("\xa0", " ") if n.strip() == "\xa0" else n)
            if t.strip() == "" and not p.runs: continue
            fmt_run(p.add_run(t), font, size, bold, color, rtl)
        elif n.tag == "br": p.add_run().add_break()
        elif n.tag == "b": inline(p, n.c, font, size, True, rtl, color)
        elif n.tag == "img": add_img(p, n)
        elif n.tag == "span":
            s = n.st; sz = px(s.get("font-size"))[0] or size; inline(p, [" "] + n.c, font, sz, bold, rtl, color)
        else: inline(p, n.c, font, size, bold, rtl, color)
BLOCK = {"div", "table"}
def render_cell(cell, td, ctx):
    cls = td.cls; rtl = "ar" in cls or "arb" in cls
    kind = ctx["cls"]; tp = "tp" in ctx["parents"]
    font = "Segoe UI" if rtl else ("Arial" if (doc_kind == "cert" or tp) else "Times New Roman")
    base = 8.5 if tp and not rtl else (9 if (tp or doc_kind == "cert") else (9.5 if rtl else 9))
    st = td.st; size = px(st.get("font-size"))[0] or base
    bold = "b" in cls or "arb" in cls or ctx.get("bold", False); color = "FFFFFF" if ctx.get("white") else None
    if "arb" in cls and not st.get("font-size"): size = 10
    align = "center" if ("c" in cls or "center" in (st.get("text-align") or "")) else ("right" if (rtl or "right" in (st.get("text-align") or "")) else None)
    first = True; buf = []
    def flush():
        nonlocal first
        if not buf: return
        p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
        para_fmt(p, rtl, align); inline(p, buf, font, size, bold, rtl, color); buf.clear()
    for ch in td.c:
        if not isinstance(ch, str) and ch.tag == "div":
            flush(); p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
            dst = ch.st; msz = px(dst.get("margin-top"))[0] or 0; dsz = px(dst.get("font-size"))[0] or size
            al = "center" if "center" in (dst.get("text-align") or "") else align
            para_fmt(p, rtl, al, before=msz * 2.83); inline(p, ch.c, font, dsz, bold or "b" in ch.cls, rtl, color)
        else: buf.append(ch)
    flush()
    if first and not cell.paragraphs[0].runs: para_fmt(cell.paragraphs[0])
def render_table(d, tnode, parents):
    cls = tnode.cls; bordered = "g" in cls or "tp" in parents
    rows = []; pending = []; colw = []
    for ch in tnode.c:
        if isinstance(ch, str):
            if ch.strip(): pending.append(ch.strip())
        elif ch.tag == "colgroup": colw = [px(c.st.get("width"))[0] for c in ch.c if not isinstance(c, str)]
        elif ch.tag == "tr": rows.append((ch, pending)); pending = []
    tail = pending
    if not rows: return
    ncols = max(sum(int(td.a.get("colspan", 1)) for td in r.c if not isinstance(td, str) and td.tag == "td") for r, _ in rows)
    if len(colw) != ncols: colw = [100.0 / ncols] * ncols
    tw = 186.0; tbl = d.add_table(rows=0, cols=ncols); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
    tblPr = tbl._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay); borders(tbl, bordered)
    for ri, (tr, pre) in enumerate(rows):
        row = tbl.add_row(); tds = [x for x in tr.c if not isinstance(x, str) and x.tag == "td"]
        ci = 0; ctxrow = {"cls": cls, "parents": parents, "white": "bh" in tr.cls or "bar" in tr.cls, "bold": "bh" in tr.cls or "bar" in tr.cls}
        for k, td in enumerate(tds):
            span = int(td.a.get("colspan", 1)); cell = row.cells[ci]
            if span > 1: cell = cell.merge(row.cells[ci + span - 1])
            w = sum(colw[ci:ci + span]) * tw / 100.0; cell.width = Mm(w)
            if k == 0 and pre: td.c.insert(0, "".join(pre) + " ")
            if ri == len(rows) - 1 and k == len(tds) - 1 and tail: td.c.append(" " + "".join(tail))
            render_cell(cell, td, ctxrow)
            if bordered: cell_margins(cell, 1.2, 1.6)
            if ctxrow["white"] and "bar" in tr.cls: shade(cell, "000000")
            elif ctxrow["white"]: shade(cell, "2F7CA3")
            h = px(td.st.get("height"))[0]
            if h or (bordered and doc_kind == "cert" and "g" in cls): row.height = Mm(h or 5.8); row.height_rule = 1
            ci += span
    return tbl
def spacer(d, pts=4):
    p = d.add_paragraph(); para_fmt(p); r = p.add_run(" "); r.font.size = Pt(pts); p.paragraph_format.line_spacing = Pt(pts)
def pagebreak(d): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def walk(d, nodes, parents=()):
    for n in nodes:
        if isinstance(n, str):
            if n.strip(): p = d.add_paragraph(); para_fmt(p); inline(p, [n.strip()], "Arial", 9, False, False)
            continue
        c = n.cls
        if n.tag == "table": render_table(d, n, parents)
        elif n.tag == "img": p = d.add_paragraph(); para_fmt(p); add_img(p, n)
        elif n.tag == "div":
            if "gap" in c or "sp" in c: spacer(d, 4); continue
            if "foot" in c: p = d.add_paragraph(); para_fmt(p); [add_img(p, i, width_mm=186) for i in n.c if not isinstance(i, str) and i.tag == "img"]; continue
            if "pg" in c:
                pagebreak(d); p = d.paragraphs[-1]; p2 = d.add_paragraph(); para_fmt(p2, align="center")
                [add_img(p2, i, height_mm=255) for i in n.c if not isinstance(i, str) and i.tag == "img"]; continue
            if "tp" in c: pagebreak(d); walk(d, n.c, parents + ("tp",)); continue
            if "h3" in c: p = d.add_paragraph(); para_fmt(p, before=4, after=2); inline(p, n.c, "Arial" if n.st.get("font-family") else "Times New Roman", 10, True, False); continue
            p = d.add_paragraph(); para_fmt(p); inline(p, n.c, "Arial", 9, False, False)
        else: walk(d, n.c, parents)
def footer_fields(d):
    p = d.sections[0].footer.paragraphs[0]; para_fmt(p, align="right")
    def fld(code):
        r = p.add_run(); fmt_run(r, "Arial", 11)
        for t, txt in (("begin", None), (None, code), ("separate", None), ("end", None)):
            if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
            else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    fmt_run(p.add_run("Page "), "Arial", 11); fld("PAGE"); fmt_run(p.add_run(" of "), "Arial", 11); fld("NUMPAGES")
ORDER = {
 "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
 "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
 "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
 "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath"}
def order_xml(d):
    parts = [d.element, d.sections[0].footer._element]
    for root in parts:
        for tag, seq in ORDER.items():
            rank = {n: i for i, n in enumerate(seq.split())}
            for el in root.iter(qn("w:" + tag)):
                kids = list(el); seen = {}
                for k in kids:
                    seen[k.tag] = k            # duplicates: keep the last one
                uniq = [k for k in kids if seen[k.tag] is k]
                for k in kids: el.remove(k)
                for k in sorted(uniq, key=lambda k: rank.get(k.tag.split("}")[1], 99)): el.append(k)
def convert(src, dst):
    global doc_kind
    raw = open(src, encoding="utf8").read(); doc_kind = "cert" if "_POLICY_" in src else "app"
    pr = P(); pr.feed(raw); d = Document()
    sec = d.sections[0]; sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.left_margin = sec.right_margin = Mm(12); sec.top_margin = Mm(7 if doc_kind == "app" else 10); sec.bottom_margin = Mm(14); sec.footer_distance = Mm(5)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
    # remove the empty first paragraph spacing: Document() starts empty
    # header table styling
    body = pr.root.c
    # title blocks inside .hd tables need their own treatment
    def hd_fix(nodes):
        for n in nodes:
            if isinstance(n, str): continue
            if n.tag == "table" and "hd" in n.cls:
                for tr in n.c:
                    if isinstance(tr, str) or tr.tag != "tr": continue
                    for td in tr.c:
                        if isinstance(td, str): continue
                        for dv in td.c:
                            if not isinstance(dv, str) and dv.tag == "div": dv.a["class"] = ((dv.a.get("class") or "") + " hdbox").strip()
            hd_fix(n.c)
    hd_fix(body)
    walk2(d, body)
    if doc_kind == "app": footer_fields(d)
    order_xml(d)
    d.save(dst)
def walk2(d, body):
    # .hd header tables: left title, right logo
    for n in body:
        if not isinstance(n, str) and n.tag == "table" and "hd" in n.cls: render_header(d, n)
        else: walk(d, [n])
def render_header(d, tnode):
    tr = next(c for c in tnode.c if not isinstance(c, str) and c.tag == "tr"); tds = [t for t in tr.c if not isinstance(t, str) and t.tag == "td"]
    tbl = d.add_table(rows=1, cols=2); tbl.autofit = False; borders(tbl, False)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tbl._tbl.tblPr.append(lay)
    L, R = tbl.rows[0].cells; L.width = Mm(110); R.width = Mm(76)
    first = True
    for dv in tds[0].c:
        if isinstance(dv, str) or dv.tag != "div": continue
        p = L.paragraphs[0] if first else L.add_paragraph(); first = False
        cl = dv.cls
        if "ttl" in cl:
            para_fmt(p, before=2, after=2); inline(p, dv.c, "Arial", 12, True, False, "1F6A99")
            ppr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr")
            for e in ("top", "left", "bottom", "right"): x = OxmlElement("w:" + e); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:space"), "4"); x.set(qn("w:color"), "CCCCCC"); bd.append(x)
            ppr.append(bd)
        elif "t1" in cl: para_fmt(p, before=8); inline(p, dv.c, "Times New Roman", 17, True, False)
        elif "t2" in cl: para_fmt(p, True, "right", before=4); inline(p, dv.c, "Segoe UI", 11, True, True)
        else: para_fmt(p); inline(p, dv.c, "Arial", 9, False, False)
    pr = R.paragraphs[0]; para_fmt(pr, align="right")
    for i in tds[1].c:
        if not isinstance(i, str) and i.tag == "img": add_img(pr, i, width_mm=66)
    spacer(d, 4)
if __name__ == "__main__": convert(sys.argv[1], sys.argv[2])
