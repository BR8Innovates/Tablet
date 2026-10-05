#!/usr/bin/env python3
"""DOCX version of the specification (same content blocks): real heading styles, TOC field, repeating table header rows, page numbers."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import importlib
sc = importlib.import_module(os.environ.get('SPEC_CONTENT', 'spec_content'))
import render_spec as rs
OUT = rs.OUT; WORK = rs.WORK
NAVY = RGBColor(0x1F, 0x38, 0x64); BLUE = RGBColor(0x0B, 0x4D, 0xA2)
d = Document(); sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(1.5))
sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.6)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); st.font.size = Pt(10.5); st.paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11.5)):
    h = d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = NAVY if name != "Heading 3" else NAVY
    h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri"); h.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 0); h.paragraph_format.space_after = Pt(5)
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)
def field(par, instr):
    for t, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if t: fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
        elif txt == instr: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        else: r.text = txt
def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
def cant_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)
TOK = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
def add_runs(p, text, size=None, bold=False, italic=False, color=None):
    for part in TOK.split(str(text)):
        if not part: continue
        if part.startswith("`"): r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 10.5) - 1.5)
        elif part.startswith("**"): r = p.add_run(part[2:-2]); r.bold = True; r.font.size = Pt(size) if size else None
        else:
            r = p.add_run(part); r.bold = bold
            if size: r.font.size = Pt(size)
        r.italic = italic
        if color and not part.startswith("`"): r.font.color.rgb = color
def para(text="", size=None, bold=False, italic=False, color=None, align=None, style=None, after=None):
    p = d.add_paragraph(style=style); add_runs(p, text, size, bold, italic, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
# cover
for _ in range(9): para("")
para(rs.TITLE, 30, True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER); para(rs.SUB, 21, True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER); para(rs.TAG, 13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
for t in (rs.VERSION, rs.DATE, rs.PREP): para(t, 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
d.add_page_break()
p = para("Table of Contents", 18, True, color=NAVY)
tp = d.add_paragraph(); field(tp, 'TOC \\o "1-3" \\h \\z \\u'); para("(Right-click the table of contents and choose Update Field to refresh the page numbers.)", 8.5, italic=True, color=RGBColor(0x66, 0x66, 0x66))
first = True; pending_break = False
for b in sc.blocks:
    k = b[0]
    if k in ("h1", "h2") and not (k == "h2" and not b[2]) or (k == "h2" and b[2]): pass
    if k == "h1":
        d.add_page_break(); d.add_heading(b[1], 1)
    elif k == "h2":
        if b[2]: d.add_page_break()
        d.add_heading(b[1], 2)
    elif k == "lbl":
        if re.match(r"^\d+\.\d+\.\d+ ", b[1]): d.add_heading(b[1], 3)
        elif b[1].endswith(":") or b[1].startswith("Journey"): para(b[1], bold=True, after=2)
        else: para(b[1], 11, True, color=BLUE, after=3)
    elif k == "p": para(b[1])
    elif k == "note": para(b[1], italic=True, color=RGBColor(0x44, 0x44, 0x44))
    elif k == "bul":
        for x in b[1]:
            pp = d.add_paragraph(style="List Bullet"); add_runs(pp, x); pp.paragraph_format.space_after = Pt(1)
    elif k == "num":
        for x in b[1]:
            pp = d.add_paragraph(style="List Number"); add_runs(pp, x); pp.paragraph_format.space_after = Pt(1)
    elif k == "kv":
        pp = d.add_paragraph(); r = pp.add_run(f"{b[1]} {b[2]}"); r.bold = True; pp.paragraph_format.space_after = Pt(1)
    elif k == "code":
        t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.cell(0, 0); shade(c, "F5F5F5"); c.text = ""
        r = c.paragraphs[0].add_run(b[1]); r.font.name = "Consolas"; r.font.size = Pt(7.5); d.add_paragraph().paragraph_format.space_after = Pt(2)
    elif k == "fig":
        d.add_picture(f"{WORK}/{b[1]}.png", width=Cm(17.5)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; para(b[2], 9.5, italic=True, color=RGBColor(0x44, 0x44, 0x44), align=WD_ALIGN_PARAGRAPH.CENTER)
    elif k == "tbl":
        head, rows, widths = b[1], b[2], b[3]
        t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        total = 18.0
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9); shade(c, "DCE6F1")
        repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells; cant_split(t.rows[-1])
            for i, v in enumerate(row):
                cells[i].text = ""; add_runs(cells[i].paragraphs[0], v, 9)
                if head[i] == "Required": cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths): row.cells[i].width = Cm(total * w / 100)
        for row in t.rows:
            for c in row.cells:
                for pp in c.paragraphs: pp.paragraph_format.space_after = Pt(1)
        d.add_paragraph().paragraph_format.space_after = Pt(2)
# footer page number
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; field(fp, "PAGE")
d.core_properties.title = rs.TITLE + " — " + rs.SUB; d.core_properties.author = "Afillar"
path = OUT + "/" + rs.OUTNAME + ".docx"; d.save(path); print("saved", path)
