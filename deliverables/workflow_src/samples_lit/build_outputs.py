import json, re, os, sys, shutil, textwrap, zipfile
sys.path.insert(0, "."); import descs
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
SRC = sys.argv[1] if len(sys.argv) > 1 else "samples_all.json"
D = json.load(open(SRC)); CASES = D["cases"]; BASE = "https://portal-gw.insuremo.com/platform/api-orchestration-test/v1/flow"
for _c in CASES:
    _r = _c["response"]
    if isinstance(_r, dict) and str(_r.get("_raw", "")).startswith("PK"): _c["response"] = {"_note": "Excel file (.xlsx), returned as binary with Content-Type " + _c["ctype"], "_binary": True}
OUT = "pkg/Afillar_Samples"; shutil.rmtree("pkg", ignore_errors=True); os.makedirs(OUT + "/jsonc")
TOKEN = "eyJhbGciOiJSUzI1NiIsInR5cCI6IkpXVCJ9.sample-access-token"
def slug(s):
    m = re.match(r"^(E?)(\d+)\.\s*(.*)$", s)
    if m: return f"{m.group(1)}{int(m.group(2)):02d}_" + re.sub(r"[^A-Za-z0-9]+", "_", m.group(3)).strip("_")
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")
def pad(n): return "  " * n
def wrap_comment(text, ind):
    lines = textwrap.wrap(text, 118 - 2 * ind) or [""]
    return [f"{pad(ind)}// {l}" for l in lines]
def jsonc(obj, route, direction, ind=0, comment=True, parent=None):
    if isinstance(obj, dict):
        if not obj: return "{}"
        out = ["{"]; items = list(obj.items())
        for i, (k, v) in enumerate(items):
            dsc = "Field error code: REQUIRED, OUT_OF_RANGE, NOT_ALLOWED, UNKNOWN_VALUE, INVALID_TOTAL, INVALID_FORMAT, INVALID_CHARACTER, INVALID_TYPE, TOO_LONG, TOO_SHORT, MISMATCH." if (parent == "errors" and k == "code") else descs.describe(route, direction, k)
            if comment: out += wrap_comment(f"{k}: " + dsc, ind + 1)
            val = jsonc(v, route, direction, ind + 1, comment, k)
            out.append(f'{pad(ind + 1)}{json.dumps(k, ensure_ascii=False)}: {val}' + ("," if i < len(items) - 1 else ""))
        out.append(pad(ind) + "}"); return "\n".join(out)
    if isinstance(obj, list):
        if not obj: return "[]"
        if all(not isinstance(x, (dict, list)) for x in obj): return json.dumps(obj, ensure_ascii=False)
        out = ["["]
        for i, x in enumerate(obj):
            val = jsonc(x, route, direction, ind + 1, comment and i == 0, parent)
            out.append(pad(ind + 1) + val + ("," if i < len(obj) - 1 else ""))
        out.append(pad(ind) + "]"); return "\n".join(out)
    return json.dumps(obj, ensure_ascii=False)
def leaves(o, path=""):
    if isinstance(o, dict):
        for k, v in o.items(): yield from leaves(v, f"{path}.{k}" if path else k)
    elif isinstance(o, list):
        if all(not isinstance(x, (dict, list)) for x in o): yield path, json.dumps(o, ensure_ascii=False)
        else:
            for i, x in enumerate(o): yield from leaves(x, f"{path}[{i}]")
    else: yield path, ("" if o is None else str(o))
def last(path): return re.sub(r"\[\d+\]", "", path).split(".")[-1]
def url_of(c): return BASE + c["path"]
def query_fields(c):
    q = c["path"].split("?", 1)[1] if "?" in c["path"] else ""
    return [tuple(p.split("=", 1)) for p in q.split("&") if p]
def req_text(c):
    if c["raw_body"] is not None: return c["raw_body"]
    return json.dumps(c["body"], indent=2, ensure_ascii=False) if c["body"] is not None else ""
PURPOSE = {"get-plans": "Returns every plan of the product variant with premium and sums insured.", "proposal": "Creates a proposal (application).", "update-proposal": "Updates an open proposal (beneficiaries, payment details).",
           "issue": "Issues the policy for an open proposal.", "load-policy": "Returns a proposal (proposalNo) or an issued policy (policyNo).", "proposal-reject": "Rejects an open proposal.",
           "doc_generation": "Generates the application form or the policy certificate as a PDF.", "application_doc_generation": "Generates the application form of a proposal as a PDF.",
           "cancel-check": "Calculates the refund for a cancellation date and raises the cancellation request (maker).", "cancel-detail": "Shows the pending cancellation request to the checker.",
           "cancel-approve": "Approves (the policy is cancelled) or rejects the pending cancellation request (checker).", "list-master-table": "Lists a master / code table for drop-downs.",
           "commission-query": "Commission by policy, by product configuration, or for a period.", "generate-feedfile": "Generates the policy feed file (Excel or JSON).", "me": "Caller identity and roles.",
           "worklist": "Work queues of the staff portal.", "dashboard": "Dashboard figures of the staff portal.", "share": "Shares a quotation, proposal or policy by e-mail and / or SMS."}
# ------------------------------------------------------------------ JSONC files
index = []
for n, c in enumerate(CASES, 1):
    fol = slug(c["folder"]); d = f"{OUT}/jsonc/{fol}"; os.makedirs(d, exist_ok=True); base = f"{d}/{slug(c['name'])}"
    head = [f"// {c['name']}", f"// {c['method']} {url_of(c)}", "// Header: Content-Type: application/json" + ("" if c["name"].startswith("E14") else "; Authorization: Bearer <access token>")]
    if c["note"]: head.append("// " + c["note"])
    body = ""
    if c["body"] is not None: body = jsonc(c["body"], c["route"], "Input")
    elif c["raw_body"] is not None: body = c["raw_body"] + "  // not valid JSON on purpose"
    else:
        body = "\n".join(wrap_comment(f"{k}: " + descs.describe(c["route"], "Input", k), 0) + [f"{k} = {v}"] if False else "" for k, v in [])
        qf = query_fields(c)
        body = "\n".join(f"// {k}: {descs.describe(c['route'], 'Input', k)}\n//   {k} = {v}" for k, v in qf) or "// (no body)"
    open(base + ".request.jsonc", "w").write("\n".join(head) + "\n" + body + "\n")
    rh = [f"// HTTP {c['status']}  ({c['ctype'].split(';')[0] or 'no body'})"]
    resp = c["response"]
    open(base + ".response.jsonc", "w").write("\n".join(rh) + "\n" + jsonc(resp, c["route"], "Error" if c["status"] >= 400 else "Output") + "\n")
    c["_files"] = f"jsonc/{fol}/{slug(c['name'])}"
# ------------------------------------------------------------------ Postman collection
def desc_md(c):
    L = [f"**Purpose:** {PURPOSE.get(c['route'], '')}", "", f"**Case:** {c['name']}", f"**Result in the example:** HTTP {c['status']}", ""]
    if c["note"]: L += [c["note"], ""]
    rows = []
    if c["body"] is not None:
        for p, v in leaves(c["body"]): rows.append((p, v, descs.describe(c["route"], "Input", last(p))))
    elif c["raw_body"] is None:
        for k, v in query_fields(c): rows.append((k, v, descs.describe(c["route"], "Input", k)))
    if rows:
        L += ["**Request fields: sample value and meaning**", "", "| Field | Sample value | Meaning |", "|---|---|---|"]
        seen = set()
        for p, v, ds in rows:
            key = re.sub(r"\[\d+\]", "[]", p)
            if key in seen: continue
            seen.add(key); L.append(f"| `{p}` | `{v}` | {ds.replace('|', '/')} |")
    L += ["", "The meaning of every response field is in the file " + c["_files"] + ".response.jsonc (// comments)."]
    return "\n".join(L)
def hdrs(c):
    h = [{"key": "Content-Type", "value": "application/json"}]
    if not c["name"].startswith("E14"): h.append({"key": "Authorization", "value": "Bearer " + TOKEN})
    return h
def pm_item(c):
    u = url_of(c); q = query_fields(c)
    url = {"raw": u, "protocol": "https", "host": ["portal-gw", "insuremo", "com"], "path": [p for p in c["path"].split("?")[0].split("/") if p]}
    url["path"] = ["platform", "api-orchestration-test", "v1", "flow"] + url["path"]
    if q: url["query"] = [{"key": k, "value": v} for k, v in q]
    req = {"method": c["method"], "header": hdrs(c), "url": url, "description": desc_md(c)}
    if c["body"] is not None: req["body"] = {"mode": "raw", "raw": json.dumps(c["body"], indent=2, ensure_ascii=False), "options": {"raw": {"language": "json"}}}
    elif c["raw_body"] is not None: req["body"] = {"mode": "raw", "raw": c["raw_body"], "options": {"raw": {"language": "text"}}}
    body = json.dumps(c["response"], indent=2, ensure_ascii=False) if not (isinstance(c["response"], dict) and c["response"].get("_binary")) else c["response"]["_note"]
    ex = {"name": f"{c['status']} - sample response", "originalRequest": {k: req[k] for k in ("method", "header", "url") if k in req} | ({"body": req["body"]} if "body" in req else {}), "status": {200: "OK", 400: "Bad Request", 401: "Unauthorized", 404: "Not Found", 409: "Conflict", 422: "Unprocessable Entity", 500: "Internal Server Error"}.get(c["status"], ""), "code": c["status"],
          "_postman_previewlanguage": "json", "header": [{"key": "Content-Type", "value": c["ctype"] or "application/json"}], "body": body}
    return {"name": c["name"], "request": req, "response": [ex]}
folders = {}
for c in CASES: folders.setdefault(c["folder"], []).append(c)
col = {"info": {"name": "Afillar Insurance APIs - samples with values", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        "description": "Every request carries literal sample values (no variables). One folder per product variant (full journey: plans, proposal, enrolment, issue, load, documents, reject, cancellation approve / reject), one for the common and staff-portal APIs and one for error cases. Each request's description lists every field with its sample value and meaning; the // commented versions of the same samples are in the jsonc folder of the package.\n\nBefore sending, replace the sample access token in the Authorization header with a valid token. Proposal and policy numbers in later requests are the ones returned by the earlier requests of the same folder when the samples were recorded; re-run a folder from its first request to get fresh numbers."},
       "item": [{"name": f, "item": [pm_item(c) for c in cs]} for f, cs in folders.items()]}
json.dump(col, open(OUT + "/Afillar_Samples.postman_collection.json", "w"), indent=2, ensure_ascii=False)
# ------------------------------------------------------------------ Excel
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cases"
ws.append(["Folder", "Case", "Method", "URL", "HTTP status", "Request sample (jsonc)", "Response sample (jsonc)"])
for c in CASES: ws.append([c["folder"], c["name"], c["method"], url_of(c), c["status"], c["_files"] + ".request.jsonc", c["_files"] + ".response.jsonc"])
w2 = wb.create_sheet("Request fields"); w2.append(["Folder", "Case", "Field", "Sample value", "Meaning (required: M mandatory, O optional, C conditional)"])
for c in CASES:
    if c["body"] is not None:
        for p, v in leaves(c["body"]): w2.append([c["folder"], c["name"], p, v, descs.describe(c["route"], "Input", last(p))])
    elif c["raw_body"] is None:
        for k, v in query_fields(c): w2.append([c["folder"], c["name"], k, v, descs.describe(c["route"], "Input", k)])
w3 = wb.create_sheet("Response fields"); w3.append(["API", "Example case", "Field", "Sample value", "Meaning"])
done = set()
for c in CASES:
    if c["status"] != 200 or c["route"] in done or (isinstance(c["response"], dict) and c["response"].get("_binary")): continue
    done.add(c["route"])
    for p, v in leaves(c["response"]): w3.append([c["route"], f"{c['folder']} / {c['name']}", p, v[:200], descs.describe(c["route"], "Output", last(p))])
for s in wb:
    for cell in s[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    s.freeze_panes = "A2"
    for col_, w in zip("ABCDEFG", (34, 46, 40, 42, 70, 50, 50)): s.column_dimensions[col_].width = w
    for row in s.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(OUT + "/Afillar_Samples.xlsx")
print("cases", len(CASES), "folders", len(folders), "req rows", w2.max_row - 1, "resp rows", w3.max_row - 1)
