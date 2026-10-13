import json,datetime
from collections import Counter,OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from manual import MANUAL
R=json.load(open("results.json"))
import sys
OUT=sys.argv[1]
cnt=Counter(r["status"] for r in R)
COL={"PASS":"C6EFCE","FAIL":"FFC7CE","OBSERVATION":"FFEB9C","BLOCKED":"D9D9D9","NOT EXECUTED":"DDEBF7"}
wb=Workbook(); thin=Side(style="thin",color="BBBBBB"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def head(ws,cols,widths):
    ws.append(cols)
    for i,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(1,i); x.font=Font(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor="1F3864"); x.alignment=Alignment(wrap_text=True,vertical="center"); ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"
# Summary
ws=wb.active; ws.title="Summary"
ex=[r for r in R]
ws.append(["Term Life (TL001) – UAT Test Execution Summary"]); ws["A1"].font=Font(bold=True,size=14)
ws.append(["Environment","uponlyuat-sandbox-in.insuremo.com (tenant uponlyuat), ChannelCode SOHAR"]); ws.append(["Execution date",datetime.date.today().isoformat()]); ws.append(["Scope","InsureMO back-end APIs (Postman collection) for product TL001 – Term Life with/without Travel. UI cases listed as Not Executed."]); ws.append([])
ws.append(["Status","Count","% of executed"]); 
for c in ws[ws.max_row]: c.font=Font(bold=True)
tot=len(R)
for s in ("PASS","FAIL","OBSERVATION","BLOCKED"):
    ws.append([s,cnt.get(s,0),round(100*cnt.get(s,0)/tot,1)]); ws.cell(ws.max_row,1).fill=PatternFill("solid",fgColor=COL[s])
ws.append(["Total executed",tot,100]); ws.append(["Not executed (UI/manual)",len(MANUAL),""]); ws.append([])
ws.append(["Module","Total","Pass","Fail","Observation"])
for c in ws[ws.max_row]: c.font=Font(bold=True)
mods=OrderedDict()
for r in R: mods.setdefault(r["module"],Counter())[r["status"]]+=1
for m,c in mods.items(): ws.append([m,sum(c.values()),c.get("PASS",0),c.get("FAIL",0),c.get("OBSERVATION",0)])
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=60
# Test cases & results
ws=wb.create_sheet("Test Cases & Results")
head(ws,["Test ID","Module","Test Case","Type","Priority","Expected Result","Actual Result","Status","Exec (s)"],[13,20,55,13,10,38,80,14,8])
for r in R:
    ws.append([r["id"],r["module"],r["title"],r["type"],r["priority"],r["expected"],r["actual"],r["status"],r["secs"]])
    for c in ws[ws.max_row]: c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=bd
    ws.cell(ws.max_row,8).fill=PatternFill("solid",fgColor=COL.get(r["status"],"FFFFFF"))
for m in MANUAL:
    ws.append([m[0],m[1],m[2],"UI / Manual","-","As per spec / Policy Lifecycle sheet","Not executed – requires Uponly UI ("+m[3]+")","NOT EXECUTED",""])
    for c in ws[ws.max_row]: c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=bd
    ws.cell(ws.max_row,8).fill=PatternFill("solid",fgColor=COL["NOT EXECUTED"])
ws.auto_filter.ref=ws.dimensions
# Defects
D=json.load(open("defects.json"))
ws=wb.create_sheet("Defects")
head(ws,["Defect ID","Severity","Module","Title","Description / Evidence","Spec reference","Linked Tests","Recommendation"],[10,10,18,45,70,32,22,50])
for d in D:
    ws.append([d["id"],d["sev"],d["module"],d["title"],d["desc"],d["spec"],d["tests"],d["rec"]])
    for c in ws[ws.max_row]: c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=bd
    ws.cell(ws.max_row,2).fill=PatternFill("solid",fgColor={"Critical":"FF9999","High":"FFC7CE","Medium":"FFEB9C","Low":"E2EFDA"}[d["sev"]])
# Spec gaps
ws=wb.create_sheet("Spec Clarifications")
head(ws,["#","Topic","Question / Gap for BA & Business"],[5,30,120])
for i,(a,b) in enumerate(json.load(open("clar.json")),1): ws.append([i,a,b]); [setattr(c,"alignment",Alignment(wrap_text=True,vertical="top")) for c in ws[ws.max_row]]
wb.save(OUT)
print("saved",OUT,cnt)
