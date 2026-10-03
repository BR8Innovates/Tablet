# ---- Diagnostics split by product (exec'd inside build_deck_b.py namespace, before housekeeping) ----
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Pt as _Pt
from pptx.dml.color import RGBColor as _RGB
from deck_metrics import STATE, IDV, J as _J

def find_slide(title):
    for i, sl_ in enumerate(prs.slides):
        t = [s_ for s_ in sl_.shapes if s_.has_text_frame and s_.text_frame.text.strip()]
        if t and t[0].text_frame.text.strip().startswith(title):
            return i
    raise KeyError(title)

def regs_p(q, prod):
    return set(q[(q['product'] == prod) & (q.reg != '')].reg)

def newrep_p(prod):
    rj = regs_p(_J, prod); ra = regs_p(AD['q'], prod); rs_ = regs_p(SD['q'], prod)
    return dict(na=len(ra - rj), ra=len(ra & rj), ns=len(rs_ - rj - ra), rs=len(rs_ & (rj | ra)), cumA=len(rj | ra), cumS=len(rj | ra | rs_))

def segs_p(d, idvd, prod):
    q = d['q']; qq = q[(q.reg != '') & (q['product'] == prod)]
    first = qq.sort_values('time').groupby('reg').agg(vage=('vage', 'first'), rto=('rto', 'first'), prov=('provider', 'first'))
    first['state'] = first.rto.astype(str).str[:2].map(STATE).fillna('Other')
    first['age'] = pd.cut(first.vage.astype(float), [-1, 3, 7, 200], labels=['0-3', '4-7', '8+']).astype(object).fillna('Unknown')
    mi = {r_: min(v for _, v in l_) for r_, l_ in idvd.items()}
    first['idv'] = pd.Series(mi)
    first['idvb'] = pd.cut(first.idv, [-1, 1e5, 3e5, 5e5, 1e6, 1e12], labels=['<1L', '1-3L', '3-5L', '5-10L', '10L+'], right=False).astype(object).fillna('Unpriced')
    pol = d['pol']; pol = pol[pol['product'] == prod].merge(first[['age', 'idvb', 'state']], left_on='reg', right_index=True, how='left')
    def seg(col, keys=None):
        l_ = first.groupby(col).size(); pc = pol.groupby(col).size(); g_ = pol.groupby(col).premium.sum()
        keys = keys or list(l_.sort_values(ascending=False).index)
        return [(k, int(l_.get(k, 0)), int(pc.get(k, 0)), int(round(g_.get(k, 0)))) for k in keys]
    carr = []
    for p_ in P3:
        gq = qq[qq.provider == p_]; pp_ = pol[pol.provider == p_]
        carr.append((p_, int(gq.reg.nunique()), len(pp_), int(round(pp_.premium.sum()))))
    return dict(carr=carr, age=seg('age'), idv=seg('idvb', ['<1L', '1-3L', '3-5L', '5-10L', '10L+']),
                state=sorted(seg('state'), key=lambda x: -x[1])[:8])

def fail_p(d, prod):
    f = d['fail']; f = f[f['product'] == prod]
    out = {}
    for p_ in P3:
        g = f[f.provider == p_]; t = g.groupby('err_n').size().sort_values(ascending=False)
        out[p_] = dict(n=len(g), calls=int(((d['q'].provider == p_) & (d['q']['product'] == prod)).sum()),
                       reasons=[(e, int(n_), owner(g[g.err_n == e].cat.iloc[0])) for e, n_ in t.head(5).items()],
                       owner={k: int((g.owner == k).sum()) for k in ['Carrier', 'Fibe', 'Fibe Mapping', 'Unclassified']})
    return out

PN = {'2W': 'Two-Wheeler (2W)', '4W': 'Four-Wheeler (4W)'}

# ---------------------------------------------------------------- Segment mix (clone)
def seg_slide(prod, at):
    src = find_slide('Segment Mix')
    ns_ = duplicate_slide(prs, src, at)
    x = lambda k: ns_.shapes[k]
    sg = segs_p(SD, IDV['2026-09'], prod)
    settext(x(0), 'Segment Mix — %s' % PN[prod]); settext(x(1), 'Who bought %s in September, and where the premium came from' % prod)
    settext(x(5), 'Carrier')
    def rows_(lst, first='Segment'):
        return [[first if first != 'Segment' else None, None, None, None]] + [[k, ind(l_), ind(po), ind(g_)] for k, l_, po, g_ in lst]
    fill_table(x(6).table, rows_(sg['carr'], 'Carrier')); fill_table(x(8).table, rows_(sg['age']))
    fill_table(x(10).table, rows_(sg['idv'])); fill_table(x(12).table, rows_(sg['state']))
    _pp = SD['pol'][SD['pol']['product'] == prod]; pols = len(_pp); gw = int(round(_pp.premium.sum()))
    best_age = max(sg['age'], key=lambda z: z[2]); top_state = sg['state'][0]
    settext(x(15), ['▪   %s: %s distinct leads, %d policies, INR %s GWP in September' % (prod, ind(sum(a_[1] for a_ in sg['age'])), pols, ind(gw)),
                    '▪   INR %s average premium on %s policies' % (ind(gw / pols) if pols else '0', prod),
                    '▪   Policies by vehicle age: %s years led with %d (INR %s)' % (best_age[0], best_age[2], ind(best_age[3])),
                    '▪   Top state by leads: %s (%d leads, %d policies)' % (top_state[0], top_state[1], top_state[2])])
    settext(x(16), 'September 2026, full month, %s only. Leads are distinct registrations quoted; the carrier table counts a vehicle once per carrier it was quoted with.' % prod)

# ---------------------------------------------------------------- Failure reasons (clone)
def fail_slide(prod, at):
    src = find_slide('Failure Reasons by Carrier')
    ns_ = duplicate_slide(prs, src, at)
    x = lambda k: ns_.shapes[k]
    fp = fail_p(SD, prod)
    settext(x(0), 'Failure Reasons by Carrier — %s' % prod)
    settext(x(1), 'The %s failures on %s quote calls, split by who is being asked to price' % (ind(sum(fp[p_]['n'] for p_ in P3)), prod))
    for p_, (hv, hs, tb, own) in {il_: (7, 8, 9, [(12, 13), (15, 16), (18, 19), (21, 22)]), dg_: (25, 26, 27, [(30, 31), (33, 34), (36, 37), (39, 40)]),
                                  ta_: (43, 44, 45, [(48, 49), (51, 52), (54, 55), (57, 58)])}.items():
        settext(x(hv), ind(fp[p_]['n'])); settext(x(hs), 'failed calls  ·  %.1f%% of its %s quote calls' % (100 * fp[p_]['n'] / fp[p_]['calls'], prod))
        rws = [[None] * 3] + [[short(e), ind(n_), o] for e, n_, o in fp[p_]['reasons']]
        fill_table(x(tb).table, rws)
        for r_ in range(1, 6):
            set_cell_color(x(tb).table, r_, 2, OWNER_COL[rws[r_][2]])
        tot_o = max(1, sum(fp[p_]['owner'].values()))
        for (vi, pi), nm_ in zip(own, ['Carrier', 'Fibe', 'Fibe Mapping', 'Unclassified']):
            settext(x(vi), ind(fp[p_]['owner'][nm_])); settext(x(pi), '%.0f%%' % (100 * fp[p_]['owner'][nm_] / tot_o))
    ih_ = SD['fail'][(SD['fail']['product'] == prod) & SD['fail'].err_n.str.contains('iHub', na=False)]
    settext(x(60), ('Platform incident   ', '%s of the %s failures are iHub integration faults (%s on 21–22 Sept) — all engineering-fixable. Owner split is by root-cause category.' % (
        ind(len(ih_)), prod, ind(int(ih_.date.isin(['2026-09-21', '2026-09-22']).sum())))))

# ---------------------------------------------------------------- New vs repeat (clone; chart rebuilt)
def newrep_slide(prod, at):
    src = find_slide('New versus Repeat Registrations')
    ns_ = duplicate_slide(prs, src, at)
    S_ = list(ns_.shapes)
    x = lambda k: S_[k]
    n_ = newrep_p(prod)
    settext(x(0), 'New versus Repeat — %s' % PN[prod]); settext(x(1), 'September\'s %s leads: how many were new vehicles versus returning' % prod)
    settext(x(6), ind(n_['ns'])); settext(x(7), 'New %s vehicles in September'); settext(x(7), 'New %s vehicles in September' % prod)
    TCx = lambda k, t, good: (settext(x(k), t), paint(x(k), GREEN if good else RED))
    TCx(8, '%s vs August' % chg(n_['ns'], n_['na']), n_['ns'] >= n_['na'])
    settext(x(10), ind(n_['rs'])); settext(x(11), 'Returning vehicles'); TCx(12, '%s vs August' % chg(n_['rs'], n_['ra']), True)
    settext(x(14), '%.1f%%' % (100 * n_['rs'] / (n_['ns'] + n_['rs']))); settext(x(15), 'Repeat share'); settext(x(16), 'August %.1f%%' % (100 * n_['ra'] / (n_['na'] + n_['ra'])))
    settext(x(18), ind(n_['cumS'])); settext(x(19), 'Cumulative unique %s vehicles' % prod); settext(x(20), '%s at end of August' % ind(n_['cumA']))
    # replace chart
    gf = x(21); L_, T_, W_, H_ = gf.left, gf.top, gf.width, gf.height
    gf._element.getparent().remove(gf._element)
    cd = CategoryChartData(); cd.categories = ['August 2026', 'September 2026']
    cd.add_series('New this month', [n_['na'], n_['ns']]); cd.add_series('Seen in an earlier month', [n_['ra'], n_['rs']])
    ch = ns_.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, L_, T_, W_, H_, cd).chart
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = _Pt(9)
    ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = _RGB(0xD9, 0xE2, 0xF0)
    ch.value_axis.tick_labels.font.size = _Pt(9); ch.category_axis.tick_labels.font.size = _Pt(10)
    for ser, col in zip(ch.plots[0].series, [_RGB(0x1A, 0x73, 0xD9), _RGB(0xC5, 0xD3, 0xEA)]):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
    ch.plots[0].has_data_labels = True; ch.plots[0].data_labels.font.size = _Pt(9); ch.plots[0].data_labels.number_format = '#,##0'; ch.plots[0].data_labels.number_format_is_linked = False
    ch.plots[0].gap_width = 80
    settext(x(23), 'The %s picture' % prod)
    settext(x(24), ['▪   New %s vehicles %s: %s → %s' % (prod, chg(n_['ns'], n_['na'])[2:].join(['fell ', '']) if n_['ns'] < n_['na'] else 'rose ' + chg(n_['ns'], n_['na'])[2:], ind(n_['na']), ind(n_['ns'])),
                    '▪   Returning vehicles %s → %s (%s)' % (ind(n_['ra']), ind(n_['rs']), chg(n_['rs'], n_['ra'])),
                    '▪   Repeat share %.1f%% → %.1f%%' % (100 * n_['ra'] / (n_['na'] + n_['ra']), 100 * n_['rs'] / (n_['ns'] + n_['rs'])),
                    '▪   Cumulative %s base: %s → %s unique vehicles' % (prod, ind(n_['cumA']), ind(n_['cumS']))])
    settext(x(25), 'Repeat = quoted in any earlier month in the API window (from 5 July 2026); each registration is assigned to one product. August restated on the same basis.')

# build from the last diagnostics slide backwards so insert positions stay valid
i_seg = find_slide('Segment Mix'); seg_slide('2W', i_seg + 1); seg_slide('4W', i_seg + 2)
i_f = find_slide('Failure Reasons by Carrier —') if False else find_slide('Failure Reasons by Carrier'); fail_slide('2W', i_f + 1); fail_slide('4W', i_f + 2)
i_n = find_slide('New versus Repeat Registrations'); newrep_slide('2W', i_n + 1); newrep_slide('4W', i_n + 2)
