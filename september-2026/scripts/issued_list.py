import gzip,glob,json,pickle,collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
D=pickle.load(open('prep.pkl','rb'))
pol=pd.concat([D['2026-08']['pol'],D['2026-09']['pol']]).reset_index(drop=True)
regs=set(pol.reg)
veh=collections.defaultdict(dict); names={}; digit_state={}
def lj(s):
    try: return json.loads(s) if s else None
    except Exception: return None
for m in ['2026-08','2026-09']:
    for f in sorted(glob.glob(f'raw_{m}/*.jsonl.gz')):
        for l in gzip.open(f,'rt'):
            r=json.loads(l); t=r['TransactionType']
            if t not in (1001,4001,2001,2003): continue
            lg=r['TransactionSyncLogs'][0] if r['TransactionSyncLogs'] else {}
            if t==2003 and r['ProviderCode']=='DIGIT':
                res=(lj(lg.get('TransResponse')) or {}).get('Result') or {}
                if res.get('CarrierPolicyNo'): digit_state[res['CarrierPolicyNo']]=res.get('StatusCode')
            req=lj(r['TransRequest']) or {}
            try: rk=req['PolicyLobList'][0]['PolicyRiskList'][0]
            except Exception: continue
            reg=rk.get('RegistrationNo')
            if reg not in regs: continue
            p=r['ProviderCode']
            if p=='TATAAIG' and rk.get('Make'): names[reg]=(rk.get('Make'),rk.get('Model'),rk.get('Variant'))
            if t in (2001,1001,4001):
                d=veh[(p,reg)]
                for k in ('ChassisNo','EngineNo','ManufactureYear','RegistrationDate','RTOCode','FuelType','VehicleAge','Make','Model','Variant','VehicleMaincode','VehicleIDV'):
                    if rk.get(k) not in (None,'') and (k not in d or t==2001): d[k]=rk.get(k)
PN={'ICICI Lombard':'IL','Go Digit':'DIGIT','Tata AIG':'TATAAIG'}
FUEL={1:'Petrol','1':'Petrol',2:'Diesel','2':'Diesel',3:'CNG','3':'CNG',4:'LPG','4':'LPG',5:'Electric','5':'Electric'}
rows=[]
for x in pol.itertuples(index=False):
    d=veh.get((PN[x.provider],x.reg),{}); nm=names.get(x.reg)
    st=x.policy_status
    if x.provider=='Go Digit':
        issued=bool(x.policy_no) and st in('COMPLETE','COMPLETED','COMPLETE/COMPLETED')
        policy_no=x.policy_no; status=('Completed' if issued else st or '')+(' (state %s)'%digit_state.get(x.carrier_policy_no) if digit_state.get(x.carrier_policy_no) else '')
        reason='' if issued else 'Digit policy status is %s — not Completed'%st
    else:
        issued=bool(x.carrier_policy_no); policy_no=x.carrier_policy_no
        status=st or 'No status returned by carrier'; reason=''
    make,model,var=(nm if nm else (d.get('Make'),d.get('Model'),d.get('Variant')))
    rows.append(dict(Month={'2026-08':'Aug 2026','2026-09':'Sep 2026'}[x.month],Issued_Date=x.date,Insurer=x.provider,Vehicle_Type=x.product,Registration_No=x.reg,
        Make=make or '',Model=model or '',Variant=var or '',Make_Model_source='Tata request (names)' if nm else ('Request codes' if make else 'Not available'),
        Manufacture_Year=str(d.get('ManufactureYear',''))[:4],Registration_Date=d.get('RegistrationDate',''),RTO=d.get('RTOCode',''),Fuel_Code=str(d.get('FuelType','')),Vehicle_Master_Code=str(d.get('VehicleMaincode','')),
        Chassis_No=d.get('ChassisNo',''),Engine_No=d.get('EngineNo',''),PlanType=x.plantype,Policy_No=policy_no,Policy_Status=status,Premium_INR=round(float(x.premium),2),Counted=('Yes' if issued else 'No'),Reason=reason))
df=pd.DataFrame(rows).sort_values(['Month','Insurer','Issued_Date'],ascending=[False,True,True]).reset_index(drop=True)
pickle.dump(df,open('issued_list.pkl','wb'))
print(df.groupby(['Month','Insurer','Counted']).size().unstack(fill_value=0)); print(df.groupby(['Month','Vehicle_Type','Counted']).size().unstack(fill_value=0))
print('missing vehicle detail rows:',int((df.Chassis_No=='').sum()),'| make/model source:',df.Make_Model_source.value_counts().to_dict())
NAVY='1F3864'; B=Border(*(Side(style='thin',color='BFBFBF'),)*4)
wb=Workbook(); ws=wb.active; ws.title='Issued Policies'
def hdr(w,r,cols,fill=NAVY):
    for j,h in enumerate(cols,1):
        c=w.cell(row=r,column=j,value=h.replace('_',' ')); c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); c.fill=PatternFill('solid',fgColor=fill); c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); c.border=B
def put(w,df_,cols,start=2):
    for i,rw in enumerate(df_.itertuples(index=False),start):
        for j,k in enumerate(cols,1):
            v=getattr(rw,k); c=w.cell(row=i,column=j,value=v); c.font=Font(name='Arial',size=9); c.border=B
            if k=='Premium_INR': c.number_format='#,##0.00'
cols=['Month','Issued_Date','Insurer','Vehicle_Type','Registration_No','Make','Model','Variant','Manufacture_Year','Registration_Date','RTO','Fuel_Code','Vehicle_Master_Code','Chassis_No','Engine_No','PlanType','Policy_No','Policy_Status','Premium_INR']
ok=df[df.Counted=='Yes']
hdr(ws,1,cols); put(ws,ok,cols); ws.freeze_panes='F2'; ws.auto_filter.ref='A1:S%d'%(len(ok)+1)
for k,wd in zip('ABCDEFGHIJKLMNOPQRS',[10,12,14,9,14,16,18,22,10,13,7,9,14,22,18,9,24,22,12]): ws.column_dimensions[k].width=wd
ws.row_dimensions[1].height=30
w2=wb.create_sheet('Summary')
w2['A1']='Issued policies — summary'; w2['A1'].font=Font(name='Arial',size=14,bold=True,color=NAVY)
w2['A2']=('Issued = ICICI Lombard and Tata AIG: a CarrierPolicyNo is present. Go Digit: a PolicyNo is present AND the policy status is Completed (COMPLETE/COMPLETED; carrier state ACTIVE). '
          'Digit policies referred to underwriting (UW_REFFERED) are not counted — see "Not counted". Vehicle type is 2W / 4W.')
w2['A2'].alignment=Alignment(wrap_text=True,vertical='top'); w2.merge_cells('A2:H2'); w2.row_dimensions[2].height=48; w2['A2'].font=Font(name='Arial',size=10,color='595959')
hdr(w2,4,['Month','Insurer','2W policies','4W policies','Total issued','GWP (INR)','Not counted'])
r=5
for m in ['Sep 2026','Aug 2026']:
    for p in ['ICICI Lombard','Go Digit','Tata AIG']:
        s=ok[(ok.Month==m)&(ok.Insurer==p)]; n=df[(df.Month==m)&(df.Insurer==p)&(df.Counted=='No')]
        for j,v in enumerate([m,p,int((s.Vehicle_Type=='2W').sum()),int((s.Vehicle_Type=='4W').sum()),len(s),round(float(s.Premium_INR.sum()),2),len(n)],1):
            c=w2.cell(row=r,column=j,value=v); c.font=Font(name='Arial',size=10); c.border=B
            if j==6: c.number_format='#,##0.00'
        r+=1
    s=ok[ok.Month==m]
    for j,v in enumerate([m,'TOTAL',int((s.Vehicle_Type=='2W').sum()),int((s.Vehicle_Type=='4W').sum()),len(s),round(float(s.Premium_INR.sum()),2),int(((df.Month==m)&(df.Counted=='No')).sum())],1):
        c=w2.cell(row=r,column=j,value=v); c.font=Font(name='Arial',size=10,bold=True); c.fill=PatternFill('solid',fgColor='D9E2F3'); c.border=B
        if j==6: c.number_format='#,##0.00'
    r+=1
for k,wd in zip('ABCDEFG',[12,16,12,12,12,14,12]): w2.column_dimensions[k].width=wd
w3=wb.create_sheet('Not counted'); nc=df[df.Counted=='No']; c3=cols+['Reason']
hdr(w3,1,c3,'C00000'); put(w3,nc,c3)
for k,wd in zip('ABCDEFGHIJKLMNOPQRST',[10,12,14,9,14,16,18,22,10,13,7,9,14,22,18,9,24,26,12,44]): w3.column_dimensions[k].width=wd
for w in wb: w.sheet_view.showGridLines=False
wb.save('Issued_Policies_Aug_Sep_2026.xlsx'); print('saved; counted',len(ok),'not counted',len(nc))
