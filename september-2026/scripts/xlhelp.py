"""Styling + formula helpers that mirror the August workbook look (Arial, navy headers, blue inputs)."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = '1F3864'
TOTAL_FILL = 'D9E2F3'
LABEL_FILL = 'F2F2F2'
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EXPECTED = {}   # (sheet, cell) -> expected value for every formula cell; checked after a LibreOffice recalc


def font(bold=False, color=None, size=10, italic=False):
    return Font(name='Arial', size=size, bold=bold, color=color, italic=italic)


def title(ws, text, sub, width_cols=8):
    ws['A1'] = text
    ws['A1'].font = font(True, NAVY, 16)
    ws['A2'] = sub
    ws['A2'].font = font(False, '595959', 10)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
    ws.row_dimensions[1].height = 25.5
    ws.row_dimensions[2].height = 30
    ws.sheet_view.showGridLines = False


def header(ws, row, labels, col=1, fill=NAVY, height=31.5):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = font(True, 'FFFFFF')
        c.fill = PatternFill('solid', fgColor=fill)
        c.alignment = Alignment(horizontal='left' if (i == 0 and col == 1) else 'center', vertical='center', wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height


def label(ws, row, col, text, total=False, wrap=False):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font(True)
    c.fill = PatternFill('solid', fgColor=TOTAL_FILL if total else LABEL_FILL)
    c.border = BORDER
    c.alignment = Alignment(wrap_text=wrap, vertical='center')
    return c


def text(ws, row, col, value, wrap=False, bold=False, color=None, border=True, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold, color)
    if border:
        c.border = BORDER
    c.alignment = Alignment(wrap_text=wrap, vertical='top' if wrap else 'center')
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


def num(ws, row, col, value, fmt='#,##0', total=False, blue=False, bold=False):
    """Static number (blue = hard-coded input pulled from source)."""
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold or total, '0000FF' if blue else None)
    c.number_format = fmt
    c.alignment = Alignment(horizontal='right')
    c.border = BORDER
    if total:
        c.fill = PatternFill('solid', fgColor=TOTAL_FILL)
    return c


def formula(ws, row, col, f, expected, fmt='#,##0', total=False, bold=False):
    c = ws.cell(row=row, column=col, value=f)
    c.font = font(bold or total)
    c.number_format = fmt
    c.alignment = Alignment(horizontal='right')
    c.border = BORDER
    if total:
        c.fill = PatternFill('solid', fgColor=TOTAL_FILL)
    EXPECTED[(ws.title, c.coordinate)] = expected
    return c


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def note(ws, row, text_, cols=8, height=None, bold=False, color=None):
    c = ws.cell(row=row, column=1, value=text_)
    c.font = font(bold, color)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    if height:
        ws.row_dimensions[row].height = height
    return c


def rng(sheet, col, n):
    return "'%s'!$%s$2:$%s$%d" % (sheet, col, col, n + 1)


class Ref(str):
    """A raw cell reference used as a criteria (not quoted)."""


def sumifs(sheet, n, **conds):
    """Occurrences (col F) filtered by A=insurer B=reg C=product D=plan E=error G=category.
    Criteria: plain str -> quoted literal; Ref -> cell reference."""
    colmap = {'ins': 'A', 'reg': 'B', 'prod': 'C', 'plan': 'D', 'err': 'E', 'cat': 'G'}
    parts = [rng(sheet, 'F', n)]
    for k, v in conds.items():
        parts.append(rng(sheet, colmap[k], n))
        parts.append(str(v) if isinstance(v, Ref) else '"%s"' % v.replace('"', '""'))
    return 'SUMIFS(' + ','.join(parts) + ')'


def pct_fmt():
    return '0.0%'
