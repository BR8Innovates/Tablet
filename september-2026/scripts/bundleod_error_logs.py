import pickle,sys
import pandas as pd
from extract import row_for
src=open('last10_logs.py').read()
ns=__import__('openpyxl.styles',fromlist=['x']).__dict__.copy(); ns.update(re=__import__('re'),json=__import__('json'),pd=pd)
exec(src[src.index('PII='):src.index('Q=D;')],ns)
redact,clean,hdrrow,put=ns['redact'],ns['clean'],ns['hdrrow'],ns['put']
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment
hits=pickle.load(open('bundleod_hits.pkl','rb'))
rows=[]
for e in hits:
    x=row_for(e); lg=e['TransactionSyncLogs'][0]
    assert 'bundle od' in e['TransactionSyncLogs'][0]['SyncResponse'].lower() and x['provider']=='Tata AIG' and not x['ok']
    rows.append(dict(e=e,lg=lg,reg=x['reg'],pr=x['product'],pt=x['plantype'],pc=x['plancode'],t=e['ClientRequestTime']))
df=pd.DataFrame([{k:v for k,v in r.items() if k not in('e','lg')}|{'trace':r['e']['TraceId'],'crid':r['e']['ClientRequestId']} for r in rows])
print(len(df),df.reg.nunique(),df['pr'].value_counts().to_dict(),df.pt.value_counts().to_dict())
wb=Workbook(); ws=wb.active; ws.title='Index'
ws['A1']='Tata AIG — "1182-please provide bundle od start date and bundle od end date" (1–22 Oct 2026)'; ws['A1'].font=Font(name='Arial',size=14,bold=True)
ws['A2']=('Source: InsureMO production integration adapter. All quote calls whose carrier error is the bundle OD dates error. "Vehicle Logs" has one log per unique vehicle (latest failed call); '
          '"All ClientRequestIds" lists every occurrence. Customer personal data redacted.'); ws.merge_cells('A2:F2'); ws['A2'].alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[2].height=45
hdrrow(ws,4,['Metric','Value'])
for i,(a,b) in enumerate([('Total occurrences',len(df)),('Unique vehicles',df.reg.nunique()),('2W occurrences',int((df['pr']=='2W').sum())),('4W occurrences',int((df['pr']=='4W').sum()))],5): put(ws,i,[a,b])
ws.column_dimensions['A'].width=24
w=wb.create_sheet('All ClientRequestIds'); hdrrow(w,1,['ClientRequestId','TraceId','Registration','Product','PlanType','PlanCode','Call time'])
for i,r in enumerate(df.sort_values('t').itertuples(),2): put(w,i,[r.crid,r.trace,r.reg,r.pr,r.pt,r.pc,r.t])
for k,wd in zip('ABCDEFG',[38,34,14,9,9,9,20]): w.column_dimensions[k].width=wd
w.freeze_panes='A2'
LH=['Registration','Product','PlanType','PlanCode','Attempts','ClientRequestId','TraceId','Call time','Carrier endpoint','Sync status','Fibe → InsureMO request','InsureMO → Carrier request','Carrier response','InsureMO response']
v=wb.create_sheet('Vehicle Logs'); hdrrow(v,1,LH)
att=df.groupby('reg').size()
last={}
for r in rows:
    if r['reg'] not in last or r['t']>last[r['reg']]['t']: last[r['reg']]=r
i=2
for reg,r in sorted(last.items(),key=lambda kv:-att[kv[0]]):
    e,lg=r['e'],r['lg']
    put(v,i,[reg,r['pr'],r['pt'],r['pc'],int(att[reg]),e['ClientRequestId'],e['TraceId'],e['ClientRequestTime'],lg['SyncUrl'],lg['SyncStatus'],clean(e['TransRequest']),clean(lg['SyncRequest']),clean(lg['SyncResponse']),clean(lg['TransResponse'])]); i+=1
for j,wd in enumerate([14,8,9,9,9,38,34,20,30,9,60,60,60,60],1): v.column_dimensions[chr(64+j)].width=wd
v.freeze_panes='B2'
wb.save('out10/Tata_BundleOD_Error_ClientRequestIds_and_Logs_Oct2026.xlsx'); print('saved',i-2,'vehicle logs')
