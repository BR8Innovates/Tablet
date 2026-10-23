import gzip,json,glob
import openpyxl
from openpyxl import Workbook
from extract import row_for
src=open('last10_logs.py').read()
import pandas as pd
ns=__import__('openpyxl.styles',fromlist=['x']).__dict__.copy(); ns.update(re=__import__('re'),json=json,pd=pd)
exec(src[src.index('PII='):src.index('Q=D;')],ns); hdrrow,put=ns['hdrrow'],ns['put']
def js(s):
    if not isinstance(s,str): return s
    try: return json.loads(s)
    except Exception: return {}
veh={}   # proposal -> vehicle dict (from 2001 create-quote)
for d in ['raw_2026-07','raw_2026-08','raw_2026-09','raw_2026-10']:
    for f in sorted(glob.glob(d+'/*.jsonl.gz')):
        for l in gzip.open(f,'rt'):
            e=json.loads(l)
            if e.get('ProviderCode')!='TATAAIG' or e['TransactionType']!=2001: continue
            x=row_for(e); p=x.get('res_proposal_no')
            if not p: continue
            risk=js(e['TransRequest'])['PolicyLobList'][0]['PolicyRiskList'][0]
            veh[p]=dict(reg=risk.get('RegistrationNo'),prod=x['product'],plan=x['plantype'],ct=risk.get('CarrierType'),main=risk.get('VehicleMaincode'),fuel=risk.get('FuelType'),
                        mfg=risk.get('ManufactureYear'),regd=risk.get('RegistrationDate'),age=risk.get('VehicleAge'),rto=risk.get('RTOCode'),chassis=risk.get('ChassisNo'),engine=risk.get('EngineNo'),
                        make=x.get('make'),model=x.get('model'),qdate=e['ClientRequestTime'][:10])
VH=['Registration No','Vehicle type','CarrierType','RTO','Registration date','Mfg year','Vehicle age (yrs)','Fuel code','Vehicle code (VehicleMaincode)','Chassis No','Engine No','Create-quote date']
def vrow(p):
    v=veh.get(p)
    if not v: return ['(not found in create-quote data)']+['']*11
    return [v['reg'],v['prod'],v['ct'],v['rto'],v['regd'],v['mfg'],v['age'],v['fuel'],v['main'],v['chassis'],v['engine'],v['qdate']]
wb=Workbook(); wb.remove(wb.active); miss=0
for title,path,nm in [('Failed Issue Policy','out10/Tata_IssuePolicy_Failures_Jul-Oct2026.xlsx','Failed Issue Policy Cases'),('Issued (CarrierPolicyNo)','out10/Tata_CarrierPolicyNo_Cases_with_Logs_Jul-Oct2026.xlsx','Tata CarrierPolicyNo Cases')]:
    w=openpyxl.load_workbook(path)[nm]; rows=list(w.iter_rows(min_row=2,values_only=True))
    fail='Failed' in title
    ws=wb.create_sheet(title)
    hdrrow(ws,1,['Date','CarrierProposalNo']+(['Insurer error','Attempts'] if fail else ['CarrierPolicyNo','Calls'])+VH+['TraceId','ClientRequestId'])
    for i,r in enumerate(rows,2):
        p=js(r[7] if fail else r[8]).get('CarrierProposalNo')
        extra=[r[2],r[11]] if fail else [r[3],r[12]]
        vr=vrow(p); miss+=vr[1]=='' 
        put(ws,i,[r[0],p]+extra+vr+[r[3] if fail else r[4], r[4] if fail else r[5]])
    for j,wd in enumerate([11,20,45,9,14,9,12,9,9,9,18,20,18,14],1): ws.column_dimensions[chr(64+j)].width=wd
    ws.freeze_panes='C2'
wb.save('out10/Tata_IssuePolicy_Cases_Vehicle_Details_Jul-Oct2026.xlsx'); print('missing',miss)
for s in wb:
    print(s.title,s.max_row-1)
    for r in list(s.iter_rows(min_row=2,max_row=4,values_only=True)): print(r[:16])
