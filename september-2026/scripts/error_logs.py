import pickle,json,re,sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
from q import q
from classify import normalise
from extract import row_for
D=pickle.load(open('prep.pkl','rb')); S=D['2026-09']
fall=S['fail']; f=fall[fall.reg!='']
PII=re.compile(r'(mobile|phone|e-?mail|first_?name|last_?name|middle_?name|full_?name|^name$|holder_?name|customer_?name|nominee|dob|date_?of_?birth|address|pan$|aadhaar|gender_?code|salutation|contact)',re.I)
def redact(o):
    if isinstance(o,dict): return {k:('[REDACTED]' if PII.search(k) and not isinstance(v,(dict,list)) and v not in ('',None) else redact(v)) for k,v in o.items()}
    if isinstance(o,list): return [redact(x) for x in o]
    return o
def clean(s):
    if s is None: return ''
    try: s2=json.dumps(redact(json.loads(s)),ensure_ascii=False)
    except Exception: s2=s
    return s2 if len(s2)<=32000 else s2[:32000]+' …[truncated]'
out=[]; index=[]
for p in ['ICICI Lombard','Go Digit','Tata AIG']:
    for pr in ['2W','4W']:
        g=f[(f.provider==p)&(f['product']==pr)]
        top=g.groupby('err_n').size().sort_values(ascending=False)
        err=top.index[0]
        ga=fall[(fall.provider==p)&(fall['product']==pr)]; tot=int((ga.err_n==err).sum()); gtot=len(ga)
        ge=g[g.err_n==err]
        veh=ge.groupby('reg').agg(att=('ok','size'),last=('time','max')).sort_values(['att','last'],ascending=[False,False])
        nveh=len(veh); pick=veh.head(10)
        index.append(dict(Insurer=p,Product=pr,Major_Error=err,Occurrences=tot,Unique_Vehicles_Affected=nveh,Vehicles_Shown=len(pick),Share=tot/gtot))
        for reg,r in pick.iterrows():
            t=ge[(ge.reg==reg)].sort_values('time').iloc[-1]
            s,d=q({'traceId':t.trace},1,5); els=d['ElementsInCurrentPage']; assert len(els)==1,(t.trace,len(els))
            e=els[0]; x=row_for(e); lg=e['TransactionSyncLogs'][0]
            assert x['provider']==p and x['product']==pr and normalise(x['err'])==err and not x['ok'],(p,pr,reg,x['err'])
            out.append(dict(Insurer=p,Product=pr,Major_Error=err,Registration_No=reg,Attempts_with_this_error=int(r.att),PlanType=t.plantype,
                            Call_Time=e['ClientRequestTime'],TraceId=e['TraceId'],ClientRequestId=e['ClientRequestId'],Carrier_Endpoint=lg['SyncUrl'],
                            Carrier_Sync_Status=lg['SyncStatus'],Sync_Start=lg['SyncStartTime'],Sync_End=lg['SyncEndTime'],
                            Fibe_to_InsureMO_Request=clean(e['TransRequest']),InsureMO_to_Carrier_Request=clean(lg['SyncRequest']),
                            Carrier_Response=clean(lg['SyncResponse']),InsureMO_Response=clean(lg['TransResponse'])))
print('rows',len(out)); pickle.dump((index,out),open('error_logs.pkl','wb'))
NAVY='1F3864'; B=Border(*(Side(style='thin',color='BFBFBF'),)*4)
wb=Workbook(); ws=wb.active; ws.title='Index'
ws['A1']='Insurer logs — major error, 10 unique vehicles per insurer and product (September 2026)'; ws['A1'].font=Font(name='Arial',size=14,bold=True,color=NAVY)
ws['A2']=('Source: InsureMO production integration adapter (queryTransactionPaged), each log re-fetched live by TraceId. For every insurer × product the single most frequent normalised error in September is shown, '
          'with the 10 unique vehicles that hit it most often (latest failed call per vehicle). Customer personal data (names, phone, email, DOB, address, PAN) is redacted; vehicle identifiers are kept.')
ws['A2'].alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells('A2:H2'); ws.row_dimensions[2].height=62; ws['A2'].font=Font(name='Arial',size=10,color='595959')
hdr=['Insurer','Product','Major error','Occurrences (Sep)','Unique vehicles affected','Vehicles shown','Share of insurer-product failures','Logs sheet']
for j,h in enumerate(hdr,1):
    c=ws.cell(row=4,column=j,value=h); c.font=Font(name='Arial',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); c.border=B
for i,r in enumerate(index,5):
    vals=[r['Insurer'],r['Product'],r['Major_Error'],r['Occurrences'],r['Unique_Vehicles_Affected'],r['Vehicles_Shown'],r['Share'],r['Insurer']+' Logs']
    for j,v in enumerate(vals,1):
        c=ws.cell(row=i,column=j,value=v); c.font=Font(name='Arial',size=10); c.border=B; c.alignment=Alignment(wrap_text=True,vertical='top')
        if j==4 or j==5: c.number_format='#,##0'
        if j==7: c.number_format='0.0%'
for k,w in zip('ABCDEFGH',[15,9,70,14,14,10,16,22]): ws.column_dimensions[k].width=w
ws.sheet_view.showGridLines=False
cols=list(out[0].keys())
for p in ['ICICI Lombard','Go Digit','Tata AIG']:
    w=wb.create_sheet(p+' Logs'); w.sheet_view.showGridLines=False
    for j,h in enumerate([c.replace('_',' ') for c in cols],1):
        c=w.cell(row=1,column=j,value=h); c.font=Font(name='Arial',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); c.border=B
    rows=[o for o in out if o['Insurer']==p]
    for i,o in enumerate(rows,2):
        for j,k in enumerate(cols,1):
            c=w.cell(row=i,column=j,value=o[k]); c.font=Font(name='Arial',size=9); c.alignment=Alignment(vertical='top',wrap_text=(j in (3,)))
        w.row_dimensions[i].height=45
    for j,wd in enumerate([14,8,46,16,11,9,19,34,38,34,9,19,19,60,60,60,60],1): w.column_dimensions[w.cell(row=1,column=j).column_letter].width=wd
    w.freeze_panes='E2'; w.auto_filter.ref='A1:%s%d'%(w.cell(row=1,column=len(cols)).column_letter,len(rows)+1)
wb.save('Insurer_Error_Logs_September_2026.xlsx'); print('saved')
for r in index: print(r['Insurer'],r['Product'],r['Occurrences'],r['Unique_Vehicles_Affected'],'|',r['Major_Error'][:70])
