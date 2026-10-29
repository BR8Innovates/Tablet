import pickle,collections
from openpyxl import Workbook
src=open('last10_logs.py').read()
import json,pandas as pd
ns=__import__('openpyxl.styles',fromlist=['x']).__dict__.copy(); ns.update(re=__import__('re'),json=json,pd=pd)
exec(src[src.index('PII='):src.index('Q=D;')],ns); hdrrow,put=ns['hdrrow'],ns['put']
o=pickle.load(open('pay90.pkl','rb'))
o['ICICI Lombard']={k:v for k,v in o['ICICI Lombard'].items() if k.startswith('pay_')}   # epn_ ids are enquiry ids, not payments
wb=Workbook(); s=wb.active; s.title='Summary'
hdrrow(s,1,['Insurer','Window','Issuance calls with payment','Unique payment txns','Txns never issued','Txns with >1 attempt','Txns submitted on >1 day','Multi-day txns never issued','Max attempts on one txn','Max days on one txn'])
sh=wb.create_sheet('Payments across multiple days'); hdrrow(sh,1,['Insurer','Payment TransactionNo','Amount','Attempts','Distinct days','First attempt','Last attempt','Issued eventually?','Successful calls','Most common error'])
al=wb.create_sheet('All attempts (multi-day)'); hdrrow(al,1,['Insurer','Payment TransactionNo','Call time','Result','Error','TraceId','ClientRequestId'])
r=2;rs=2;ra=2
for name,g in o.items():
    days=lambda v:len({x[0][:10] for x in v})
    never=[k for k,v in g.items() if not any(x[1] for x in v)]
    md=[k for k,v in g.items() if days(v)>1]
    put(s,r,[name,'31 Jul – 29 Oct 2026',sum(len(v) for v in g.values()),len(g),len(never),sum(1 for v in g.values() if len(v)>1),len(md),sum(1 for k in md if k in never),max(len(v) for v in g.values()),max(days(v) for v in g.values())]); r+=1
    for k in sorted(md,key=lambda k:(-len(g[k]))):
        v=sorted(g[k]); er=collections.Counter(x[2] for x in v if not x[1]).most_common(1)
        put(sh,rs,[name,k,v[0][3],len(v),days(v),v[0][0],v[-1][0],'Yes' if any(x[1] for x in v) else 'NO',sum(x[1] for x in v),er[0][0] if er else '']); rs+=1
        for x in v: put(al,ra,[name,k,x[0],'Success' if x[1] else 'Failed',x[2],x[4],x[5]]); ra+=1
for w,ws_ in((s,[16,22,14,12,12,12,12,14,12,12]),(sh,[16,26,9,9,9,20,20,12,10,60]),(al,[16,26,20,9,60,34,38])):
    for j,x in enumerate(ws_,1): w.column_dimensions[chr(64+j)].width=x
wb.save('out10/All_Insurers_Payment_Resubmissions_90d.xlsx')
for row in s.iter_rows(min_row=2,values_only=True): print(row)
