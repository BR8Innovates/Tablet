import json,pickle,pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
D=pickle.load(open('prep.pkl','rb')); il=pickle.load(open('issued_list.pkl','rb'))
INS={'33':('Go Digit','DIGIT'),'25':('Tata AIG','TATAAIG')}
t=lambda s:datetime.fromisoformat(s)
def build(m,label,sysfile):
    pol=D[m]['pol']; sysd=json.load(open(sysfile)); ilm=il[il.Month==label]; stat={(r.Insurer,r.Policy_No):r.Policy_Status for r in ilm.itertuples()}; rows=[]; used=set()
    for x in sorted(sysd,key=lambda z:z['IssueDate']):
        name,code=INS[x['CarrierCode']]; cp=x.get('CarrierPolicyNo'); c=pol[pol.provider==name]
        mm=c[c.carrier_policy_no==cp] if cp else c[(abs(c.premium-x['DuePremium'])<0.5)&(c.first_time.map(lambda s:abs((t(s)-t(x['IssueDate'])).total_seconds())<=10))&(~c.carrier_policy_no.isin(used))]
        assert len(mm)==1,(x['PolicyNo'],len(mm)); mm=mm.iloc[0]; used.add(mm.carrier_policy_no)
        rows.append(dict(PolicyNo=x['PolicyNo'],InsurerCode=int(x['CarrierCode']),Insurer=code,CarrierPolicyNo=mm.carrier_policy_no,IssueDate=x['IssueDate'].replace('T',' ')[:16],Premium=round(float(x['DuePremium']),2),
            Policy_Status='Effective',Vehicle_Type=mm['product'],Registration_No=mm.reg,Source='InsureMO policy system',Carrier_Status=stat.get((name,mm.carrier_policy_no),'') if name=='Go Digit' else ''))
    for _,mm in pol[pol.provider=='ICICI Lombard'].iterrows():
        rows.append(dict(PolicyNo='',InsurerCode=10,Insurer='IL',CarrierPolicyNo=mm.carrier_policy_no,IssueDate=mm.first_time.replace('T',' ')[:16],Premium=round(float(mm.premium),2),
            Policy_Status='Issued (IL)',Vehicle_Type=mm['product'],Registration_No=mm.reg,Source='Integration adapter (IL proposal response)',Carrier_Status=mm.policy_status or ''))
    return pd.DataFrame(rows).sort_values('IssueDate').reset_index(drop=True)
S=build('2026-09','Sep 2026','policy_sys_sep.json'); A=build('2026-08','Aug 2026','policy_sys_aug.json')
res={}
for k,d in (('Sep',S),('Aug',A)):
    rem=d[d.Carrier_Status.str.startswith('UW_REFFERED')]; keep=d[~d.index.isin(rem.index)]
    res[k]=(keep.reset_index(drop=True),rem.reset_index(drop=True)); print(k,'total',len(d),'| removed UW referred',len(rem),'| list',len(keep),keep.groupby('Insurer').size().to_dict(),'| premium',round(keep.Premium.sum(),2))
pickle.dump(res,open('policy_lists.pkl','wb'))
NAVY='1F3864'; B=Border(*(Side(style='thin',color='BFBFBF'),)*4)
cols=['PolicyNo','InsurerCode','Insurer','CarrierPolicyNo','IssueDate','Premium','Policy_Status','Vehicle_Type','Registration_No','Source']
heads=['PolicyNo','InsurerCode','Insurer','CarrierPolicyNo','IssueDate','Premium','Policy Status','Vehicle Type','Registration No','Source']
wb=Workbook(); first=[True]
def sheet(name,df,extra=None,color=NAVY):
    ws=wb.active if first[0] else wb.create_sheet(); first[0]=False; ws.title=name; cs=cols+(extra or []); hs=heads+[e.replace('_',' ') for e in (extra or [])]
    for j,h in enumerate(hs,1):
        c=ws.cell(row=1,column=j,value=h); c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); c.fill=PatternFill('solid',fgColor=color); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=B
    for i,r in enumerate(df.itertuples(index=False),2):
        for j,k in enumerate(cs,1):
            c=ws.cell(row=i,column=j,value=getattr(r,k)); c.font=Font(name='Arial',size=10); c.border=B
            if k=='Premium': c.number_format='#,##0.00'
    n=len(df)+1; ws.cell(row=n+1,column=1,value='TOTAL').font=Font(name='Arial',bold=True)
    ws.cell(row=n+1,column=2,value='=COUNTA(D2:D%d)'%n).font=Font(name='Arial',bold=True)
    c=ws.cell(row=n+1,column=6,value='=SUM(F2:F%d)'%n); c.font=Font(name='Arial',bold=True); c.number_format='#,##0.00'
    for k,w in zip('ABCDEFGHIJK',[36,12,11,34,17,12,16,12,15,42,18]): ws.column_dimensions[k].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref='A1:%s%d'%('K' if extra else 'J',n); ws.sheet_view.showGridLines=False; ws.row_dimensions[1].height=28
sheet('September',res['Sep'][0]); sheet('August',res['Aug'][0])
rem=pd.concat([res['Sep'][1].assign(Month='Sep 2026'),res['Aug'][1].assign(Month='Aug 2026')])
sheet('Removed - Digit UW referred',rem,['Carrier_Status'],'C00000')
ws=wb.create_sheet('Notes'); ws.sheet_view.showGridLines=False
notes=['Policies issued — logic and sources','',
 'Go Digit and Tata AIG: InsureMO policy system (queryPolicy by IssueDate). PolicyNo = InsureMO policy number; Policy Status 2 = Effective. Tata CarrierPolicyNo is joined from the integration-adapter issue-policy response (issue time + premium; all matched uniquely).',
 'ICICI Lombard: added from the integration adapter (IL proposal response with a CarrierPolicyNo). InsureMO\'s policy system holds no ICICI policies, so PolicyNo is blank.',
 'Removed: Go Digit policies whose carrier status is UW_REFFERED (see the red sheet). Check result: these ARE platform policies — each exists in InsureMO with a POCOMMOTOR PolicyNo and Effective status, indistinguishable from a Completed Digit policy — so they are excluded here by your rule, not because they are missing from the platform.',
 'Counts after removal — September: %d policies, INR %s. August: %d policies, INR %s.'%(len(res['Sep'][0]),format(round(res['Sep'][0].Premium.sum(),2),',.2f'),len(res['Aug'][0]),format(round(res['Aug'][0].Premium.sum(),2),',.2f'))]
for i,s in enumerate(notes,1):
    c=ws.cell(row=i,column=1,value=s); c.alignment=Alignment(wrap_text=True,vertical='top'); c.font=Font(name='Arial',size=14 if i==1 else 10,bold=(i==1),color=NAVY if i==1 else None)
ws.column_dimensions['A'].width=130
for i in range(3,7): ws.row_dimensions[i].height=44
wb.save('Policies_Issued_Aug_Sep_2026_IL_added_UWreferred_removed.xlsx'); print('saved')
